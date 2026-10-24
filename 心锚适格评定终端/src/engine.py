"""Workbook import, offline recognition and deterministic recommendation scoring."""
import hashlib,json,os,re,sqlite3,subprocess,uuid
from contextlib import closing
from pathlib import Path
from decimal import Decimal
import openpyxl
from PIL import Image

VERSION='0.6.10'
IMPORT_VERSION=5
TYPES={'偶像':1,'图籍':2,'钥匙':3}
def plain(v):return '' if v is None else str(v).strip()
def segments(cell):
    v=cell.value
    def style(f):
        color=getattr(getattr(f,'color',None),'rgb',None)
        return bool(getattr(f,'b',False)),isinstance(color,str) and color.upper().endswith('FF0000')
    if v is None:return []
    if isinstance(v,str) or not hasattr(v,'__iter__'):
        b,r=style(cell.font);return [(str(v),b,r)]
    result=[]
    for run in v:
        b,r=style(cell.font if isinstance(run,str) else run.font)
        result.append((str(run),b,r))
    return result

def parse_entry(cell,kind,index):
    runs=segments(cell);raw=''.join(t for t,b,r in runs)
    if not raw.strip():return None
    parts=[];cursor=0
    for text,bold,red in runs:
        parts.append((cursor,cursor+len(text),bold,red));cursor+=len(text)
    tokens=[];ops=[]
    for m in re.finditer(r'[^>＞=＝≥]+|[>＞=＝≥]',raw):
        text=m.group().strip()
        if not text:continue
        if text in '>＞=＝≥':ops.append({'＞':'>','＝':'='}.get(text,text));continue
        overlaps=[x for x in parts if x[0]<m.end() and x[1]>m.start()]
        direction='min' if re.search(r'[（(]\s*min!?\s*[）)]',text,re.I) else 'max'
        unimportant='不重要' in text
        text=re.sub(r'[（(]\s*(?:min!?|不重要)\s*[）)]','',text,flags=re.I).strip()
        tokens.append({'t':text,'bold':any(x[2] for x in overlaps),'red':any(x[3] for x in overlaps),
                       'others':text in ('其他','其它','其余'),'direction':direction,'unimportant':unimportant})
    return {'kind':kind,'idx':index,'raw':raw.strip(),'tokens':tokens,'ops':ops,'all_equal':raw.strip() in ('皆可','不重要'),'source':cell.coordinate}

def import_workbook(path):
    data=Path(path).read_bytes();digest=hashlib.sha256(data).hexdigest()
    import io
    w=openpyxl.load_workbook(io.BytesIO(data),data_only=True,rich_text=True)
    required=['心锚百科','词条数值','六星角色心锚推荐','特性百科']
    if any(n not in w.sheetnames for n in required):raise ValueError('表格缺少必要工作表：'+', '.join(n for n in required if n not in w.sheetnames))
    rules=[];warnings=[];ids={}
    for cells in w['词条数值'].iter_rows(min_row=5):
        seq=plain(cells[0].value);label=plain(cells[1].value)
        if not seq and not label:continue
        if not re.fullmatch(r'[123]\.[1234]\.\d+',seq):raise ValueError('词条数值编号异常：'+cells[0].coordinate)
        if seq in ids:warnings.append('数值编号重复 '+seq+'，已按行号分别保存')
        ids[seq]=cells[0].row
        slot,pos,_=map(int,seq.split('.'));unit='percent' if '%' in cells[2].number_format else 'flat'
        vals=[]
        for tier,col in enumerate(range(2,14,2),1):
            c=cells[col]
            if isinstance(c.value,(float,int)):
                vals.append({'tier':tier,'value':str(Decimal(str(c.value))*(100 if unit=='percent' else 1)),'source':c.coordinate})
            elif c.value not in (None,'无'):raise ValueError('数值不是数字：'+c.coordinate)
        if not vals:raise ValueError('词条没有1级数值：'+seq)
        if pos==4 and any(x['tier']!=1 for x in vals):raise ValueError('第四词条出现低档，请核对原表')
        mature=[]
        for tier,col in enumerate(range(3,14,2),1):
            c=cells[col]
            if isinstance(c.value,(float,int)):
                mature.append({'tier':tier,'value':str(Decimal(str(c.value))*(100 if unit=='percent' else 1)),'source':c.coordinate})
            elif c.value not in (None,'无'):raise ValueError('30级数值不是数字：'+c.coordinate)
        rules.append({'id':cells[0].row,'slot':slot,'pos':pos,'stat':re.sub(r'[12]$','',label),'variant':label,'unit':unit,'values':vals,'values30':mature})
    anchors={}
    for row in w['心锚百科'].iter_rows(min_row=4):
        name=plain(row[1].value)
        if not name:continue
        typ=plain(row[2].value)
        if typ not in TYPES:raise ValueError('未知心锚类型：'+typ)
        pools={1:[plain(row[10].value)],2:[plain(row[11].value)],3:[plain(x.value) for x in row[12:15]],4:[plain(x.value) for x in row[15:18]]}
        anchors[name]={'name':name,'slot':TYPES[typ],'type':typ,'jobs':[plain(x.value) for x in row[4:10] if x.value],
                       'pools':{str(k):[x for x in v if x] for k,v in pools.items()},
                       'traits':{'1':[plain(row[18].value)] if row[18].value else [],'2':[plain(x.value) for x in row[19:23] if x.value]},
                       'source':row[1].coordinate,'origin':plain(row[3].value)}
    def parse_characters(ws):
        # Six-star and five-star sheets share one layout: a 7-row block per
        # character, recommendations in columns 6-11 (3 slots x 2 ranks).
        chars=[]
        for head in range(4,ws.max_row+1,7):
            name=plain(ws.cell(head,2).value)
            if not name:continue
            ch={'name':name,'job':plain(ws.cell(head,3).value),'signet':plain(ws.cell(head,4).value),'damage':plain(ws.cell(head,5).value),
                'note':plain(ws.cell(head+1,1).value),'recs':[]}
            for slot,cols in enumerate([(6,7),(8,9),(10,11)],1):
                for rank,col in enumerate(cols,1):
                    anchor=plain(ws.cell(head,col).value)
                    rec={'anchor':anchor,'slot':slot,'rank':rank,'source':ws.cell(head,col).coordinate,'entries':[],'active':anchor not in ('','无','观望')}
                    if rec['active']:
                        for i in range(1,7):
                            e=parse_entry(ws.cell(head+i,col),'entry' if i<=4 else 'trait',i if i<=4 else i-4)
                            if e:rec['entries'].append(e)
                        if anchor not in anchors:warnings.append('推荐心锚未在百科中找到：'+anchor)
                    ch['recs'].append(rec)
            chars.append(ch)
        return chars
    chars=parse_characters(w['六星角色心锚推荐'])
    chars5=parse_characters(w['五星角色心锚推荐']) if '五星角色心锚推荐' in w.sheetnames else []
    if not chars or not rules:raise ValueError('工作表未解析到有效数据')
    notes5=plain(w['五星角色心锚推荐']['A2'].value) if chars5 else ''
    return {'version':IMPORT_VERSION,'hash':digest,'source':str(Path(path).resolve()),'filename':Path(path).name,
            'characters':chars,'characters5':chars5,'anchors':anchors,'rules':rules,'warnings':sorted(set(warnings)),
            'notes':plain(w['词条数值']['A2'].value),'recommendation_notes':plain(w['六星角色心锚推荐']['A2'].value),
            'recommendation_notes5':notes5}

class Store:
    def __init__(self,root):
        self.root=Path(root);self.folder=self.root/'表格';self.folder.mkdir(parents=True,exist_ok=True)
        self.data=self.root/'数据';self.data.mkdir(exist_ok=True);self.db=self.data/'rules.db'
        self.config=self.data/'settings.json'
    def files(self):return sorted(p for p in self.folder.glob('*.xlsx') if not p.name.startswith('~$'))
    def cached(self):
        if not self.db.exists():return None
        with closing(sqlite3.connect(self.db.as_uri()+'?mode=ro',uri=True)) as c:return json.loads(c.execute('select payload from snapshot').fetchone()[0])
    def load(self,path):
        path=Path(path)
        digest=hashlib.sha256(path.read_bytes()).hexdigest();old=self.cached()
        if old and old.get('version')==IMPORT_VERSION and old['hash']==digest and old.get('filename')==path.name:return old,False
        model=import_workbook(path);tmp=self.data/('rules-'+uuid.uuid4().hex+'.tmp')
        try:
            with closing(sqlite3.connect(tmp)) as c:
                c.executescript('CREATE TABLE snapshot(payload TEXT NOT NULL);CREATE TABLE characters(name TEXT PRIMARY KEY,payload TEXT);CREATE TABLE value_rules(id INTEGER PRIMARY KEY,payload TEXT);CREATE TABLE anchors(name TEXT PRIMARY KEY,payload TEXT);')
                c.execute('insert into snapshot values (?)',(json.dumps(model,ensure_ascii=False),))
                c.executemany('insert into characters values (?,?)',[(x['name'],json.dumps(x,ensure_ascii=False)) for x in model['characters']])
                c.executemany('insert into value_rules values (?,?)',[(x['id'],json.dumps(x,ensure_ascii=False)) for x in model['rules']])
                c.executemany('insert into anchors values (?,?)',[(n,json.dumps(x,ensure_ascii=False)) for n,x in model['anchors'].items()])
                assert c.execute('pragma integrity_check').fetchone()[0]=='ok'
                c.commit()
            os.replace(tmp,self.db)
        finally:
            if tmp.exists():tmp.unlink()
        return model,True

def classify(model,slot,pos,stat,value,unit,star=None):
    try:v=Decimal(str(value))
    except Exception:return {'tiers':[],'state':'review','text':'数值缺失或无效'}
    if not v.is_finite() or v<0:return {'tiers':[],'state':'review','text':'数值异常'}
    hits=[];tol=Decimal('.01') if unit=='percent' else Decimal(0)
    for r in model['rules']:
        if (r['slot'],r['pos'],r['stat'],r['unit'])!=(slot,pos,stat,unit):continue
        for x in r['values']:
            if abs(Decimal(x['value'])-v)<=tol and (star is None or (x['tier']==1)==star):hits.append({'tier':x['tier'],'variant':r['variant'],'source':x['source']})
    tiers=sorted(set(h['tier'] for h in hits))
    state='review' if not tiers or min(tiers)<=2<max(tiers) else 'top2' if max(tiers)<=2 else 'low'
    text=' / '.join(str(t)+'档' for t in tiers) if tiers else '数值/单位/星标冲突或规则未覆盖'
    return {'tiers':tiers,'state':state,'text':text,'candidates':hits}

def preference(entry,actual,legal):
    if actual is None:return 0,False
    if entry['all_equal']:return (1,False) if actual in legal else (0,False)
    p=1;explicit={x['t'] for x in entry['tokens'] if not x['others']}
    for i,t in enumerate(entry['tokens']):
        if actual==t['t'] or (t['others'] and actual in legal and actual not in explicit):return p,t['bold'] or t['red']
        if i<len(entry['ops']):p=max(0,p-{'>':.2,'≥':.1,'=':0}[entry['ops'][i]])
    return 0,False

def value_quality(model,slot,pos,stat,value,unit):
    try:v=Decimal(str(value))
    except Exception:return {'scores':[],'basis':'数字无法读取'}
    if not v.is_finite() or v<=0:return {'scores':[],'basis':'数值异常，需复核截图'}
    rules=[r for r in model['rules'] if (r['slot'],r['pos'],r['stat'],r['unit'])==(slot,pos,stat,unit)]
    if not rules:return {'scores':[],'basis':'表格没有该位置、属性和单位的满值规则'}
    if pos==4:return {'scores':[1.0],'basis':'按表格说明：第四词条存在即满档'}
    tol=Decimal('.01') if unit=='percent' else Decimal(0)
    candidates=[]
    for r in rules:
        maximum=next(Decimal(x['value']) for x in r['values'] if x['tier']==1)
        minimum=min(Decimal(x['value']) for x in r['values'])
        if minimum-tol<=v<=maximum+tol:
            candidates.append(min(1.0,float(v/maximum)))
    if candidates:return {'scores':candidates,'basis':'数值在表内范围，按当前值/满值评分；不强行判定档位'}
    ranges='、'.join(str(min(Decimal(x['value']) for x in r['values']))+'—'+str(next(Decimal(x['value']) for x in r['values'] if x['tier']==1)) for r in rules)
    return {'scores':[],'basis':'超出表内范围 '+ranges+('%' if unit=='percent' else '')+'，需复核截图或更新表格'}

def normalize_stat(label,value,unit,legal):
    raw_stat,raw_value=label,value
    aliases={'受到治疗':'受治疗','全部易伤':'全部免伤','物质易伤':'物质免伤','异能易伤':'异能免伤'}
    canonical=aliases.get(label)
    if canonical in legal:
        if label.endswith('易伤'):
            # A reduction in vulnerability is the source workbook's damage immunity.
            # Never strip the sign from arbitrary negative stats or positive vulnerability.
            if unit=='percent' and Decimal(value)<0:label=canonical;value=str(-Decimal(value))
        else:label=canonical
    return {'stat':label,'value':value,'unit':unit,'star':None,'raw_stat':raw_stat,'raw_value':raw_value}

def anchor_recommendations(model,name):
    """Recommending characters for one anchor as (character, rec) pairs.
    Six-star rows come first so a caller can show them with priority; the
    five-star block is returned separately and stays advisory only."""
    six=[(c,r) for c in model.get('characters') or [] for r in c['recs'] if r.get('active') and r.get('anchor')==name]
    five=[(c,r) for c in model.get('characters5') or [] for r in c['recs'] if r.get('active') and r.get('anchor')==name]
    return six,five

def catalog_anchors(model):
    """Every anchor the encyclopedia can open: the workbook list plus any
    recommended name missing from it (the import already warns about those)."""
    names=set(model.get('anchors') or {})
    for c in (model.get('characters') or [])+(model.get('characters5') or []):
        for r in c['recs']:
            if r.get('active') and r.get('anchor'):names.add(r['anchor'])
    return sorted(names)

def catalog_people(model):
    """Every character the catalog can open: six-star rows first so they keep
    priority, five-star rows after. Each entry keeps the parse_characters
    shape and gains 'tier' (6 or 5) so a caller can filter or label the
    source table without re-deriving it. Copies are deep: the catalog only
    reads, but a stray edit must never reach the shared model."""
    import copy
    people=[dict(copy.deepcopy(c),tier=6) for c in model.get('characters') or []]
    people+=[dict(copy.deepcopy(c),tier=5) for c in model.get('characters5') or []]
    return people

def evaluate(model,obs,selected=None,policy=None):
    from scoring import evaluate_dual
    return evaluate_dual(model,obs,selected,policy)


def recognize(root,path,model,stop=None):
    root=Path(root);out=root/'数据'/'识别记录'/uuid.uuid4().hex;out.mkdir(parents=True)
    with Image.open(path) as img:
        img=img.convert('RGB');w,h=img.size
        if abs(w/h-16/9)>.06:raise ValueError('首版支持接近16:9的完整游戏截图，请先裁掉窗口边框和黑边')
        img=img.resize((1280,720),Image.Resampling.LANCZOS);img.save(out/'input.png')
        from vision import rarity
        stars=rarity(img)
        if stars in (1,2):
            # A low-star anchor can still be a usable event item. Read the name
            # only: its entries and values are fixed by the workbook and are
            # never scored, so the zoom crops are skipped for speed.
            label='一星' if stars==1 else '二星'
            from ocr_worker import worker
            worker(root).request(out/'input.png',out/'ocr.json',stop)
            raw=json.loads((out/'ocr.json').read_text(encoding='utf-8'))
            items=[]
            for block in raw:
                for x in (block.get('detail') or {}).get('all',[]):
                    if 'box' in x and x.get('text'):items.append(x)
            anchor,name,name_issues=parse_name(items,model,fuzzy=False)
            def band(lo,hi):return sorted([x for x in items if lo<=x['box'][1]+x['box'][3]/2<=hi],key=lambda x:x['box'][0])
            lv=''.join(x['text'].replace(' ','') for x in band(108,148));m=re.search(r'(\d+)\s*[/／]\s*\d+',lv)
            level=int(m.group(1)) if m else None
            known=anchor in (model.get('anchors') or {})
            obs={'anchor':anchor or label+'心锚','level':level,'rarity':stars,'entries':{},'traits':{},'issues':name_issues,
                 'image':str(out/'input.png'),'original':str(path),'evidence':str(out),'raw_ocr':items,
                 'source':(label+'活动心锚；词条数值按原表固定，未做词条OCR' if known else label+'心锚；未做词条OCR')}
            (out/'observation.json').write_text(json.dumps(obs,ensure_ascii=False,indent=2),encoding='utf-8')
            return obs
        img.crop((470,390,520,490)).resize((200,400),Image.Resampling.LANCZOS).save(out/'input.png.zoom.png')
        img.crop((335,155,530,335)).resize((780,720),Image.Resampling.LANCZOS).save(out/'input.png.rows.png')
    app=root/'程序文件'
    from ocr_worker import worker
    worker(root).request(out/'input.png',out/'ocr.json',stop)
    raw=json.loads((out/'ocr.json').read_text(encoding='utf-8'))
    items=[]
    for block in raw:
        for x in (block.get('detail') or {}).get('all',[]):
            if 'box' in x and x.get('text'):
                if block.get('zoom'):
                    x=dict(x,box=[470+x['box'][0]/4,390+x['box'][1]/4,x['box'][2]/4,x['box'][3]/4],zoom=True)
                if block.get('rows_zoom'):
                    x=dict(x,box=[335+x['box'][0]/4,155+x['box'][1]/4,x['box'][2]/4,x['box'][3]/4],rows_zoom=True)
                items.append(x)
    obs=parse_ocr(items,model)
    from vision import star,empty_row,blank_trait
    obs['presence']={};obs['trait_presence']={}
    for key,item in obs['entries'].items():
        if item:
            item['star']=star(img,int(key));obs['presence'][key]='present'
        elif int(key)>=3 and empty_row(img,int(key)):
            obs['presence'][key]='absent'
            obs['issues']=[s for s in obs['issues'] if s!='词条'+key+'未识别，需确认是否为空']
        else:obs['presence'][key]='unknown'
    for key,item in obs['traits'].items():
        obs['trait_presence'][key]='present' if item else 'absent' if blank_trait(img,int(key)) else 'unknown'
        if obs['trait_presence'][key]=='absent':obs['issues']=[s for s in obs['issues'] if s!='特性'+key+'未识别，需确认是否为空']
    obs['source']='本地自动识别'
    obs.update({'image':str(out/'input.png'),'original':str(path),'evidence':str(out),'raw_ocr':items})
    (out/'observation.json').write_text(json.dumps(obs,ensure_ascii=False,indent=2),encoding='utf-8')
    return obs

def parse_name(items,model,fuzzy=True):
    """Read the anchor name from the base OCR pass. The leading item icon is
    low-confidence noise; simplified/traditional OCR variants are resolved
    against the workbook, and an ambiguous name is never silently mismatched.
    The fuzzy fallback must stay off for one/two-star reads: most low-star
    anchors are ordinary items missing from the workbook, and a near match
    would silently turn them into event anchors (e.g. 爱伦的家庭 scored 0.6
    against 爱伦的葬礼). A miss is harmless - the item is skipped - while a
    wrong match mislabels it, so low-star names match exactly or not at all."""
    anchors=model.get('anchors') or {}
    def band(lo,hi):return sorted([x for x in items if not x.get('zoom') and not x.get('rows_zoom') and lo<=x['box'][1]+x['box'][3]/2<=hi],key=lambda x:x['box'][0])
    def text(xs):return ''.join(x['text'].replace(' ','') for x in xs)
    name_band=band(60,105);i=0
    while i<len(name_band) and float(name_band[i].get('score',0))<.5:i+=1
    name=text(name_band[i:] if i and i<len(name_band) else name_band)
    issues=[];found=[a for a in anchors if a in name]
    if len(found)!=1 and name and fuzzy:
        import difflib
        ranked=sorted(((difflib.SequenceMatcher(None,a,name).ratio(),a) for a in anchors),reverse=True)
        if ranked and ranked[0][0]>=.6 and (len(ranked)<2 or ranked[0][0]-ranked[1][0]>=.1):found=[ranked[0][1]]
    anchor=found[0] if len(found)==1 else name
    if len(found)!=1:issues.append('名称需核对：'+name)
    if any(float(x.get('score',0))<.8 for x in band(60,105) if x['box'][0]>=340):issues.append('名称OCR置信度偏低')
    return anchor,name,issues

def parse_ocr(items,model):
    def band(lo,hi):return sorted([x for x in items if not x.get('zoom') and not x.get('rows_zoom') and lo<=x['box'][1]+x['box'][3]/2<=hi],key=lambda x:x['box'][0])
    def text(xs):return ''.join(x['text'].replace(' ','') for x in xs)
    anchor,name,name_issues=parse_name(items,model)
    issues=list(name_issues)
    lv=text(band(108,148));m=re.search(r'(\d+)\s*[/／]\s*\d+',lv);level=int(m.group(1)) if m else None
    if not m:issues.append('等级未识别')
    entries={};traits={}
    for i,(lo,hi) in enumerate([(162,204),(205,247),(248,287),(288,330)],1):
        legal=model['anchors'].get(anchor,{}).get('pools',{}).get(str(i),[])
        original=band(lo,hi)
        enlarged=sorted([x for x in items if x.get('rows_zoom') and lo<=x['box'][1]+x['box'][3]/2<=hi],key=lambda x:x['box'][0])
        def candidate(xs):
            s=text(xs).replace('％','%').replace('，','.').replace('。','.')
            m=re.search(r'(-?[0-9]+(?:\.[0-9]+)?)(%)?$',s)
            if not xs or not m:return None
            label=s[:m.start()].strip(' +:：');unit='percent' if m.group(2) else 'flat'
            if unit=='percent' and label+'%' in legal:label+='%'
            return normalize_stat(label,m.group(1),unit,legal)
        a=candidate(original);z=candidate(enlarged)
        def good(c,xs):return c and c['stat'] in legal and Decimal(c['value'])>0 and all(float(x.get('score',0))>=.8 for x in xs)
        if good(z,enlarged):
            xs=enlarged
            if good(a,original) and (a['stat'],a['value'],a['unit'])!=(z['stat'],z['value'],z['unit']):
                issues.append('词条'+str(i)+'两次OCR结果不一致')
        else:xs=original or enlarged
        s=text(xs).replace('％','%').replace('，','.').replace('。','.')
        if not xs:entries[str(i)]=None;issues.append('词条'+str(i)+'未识别，需确认是否为空');continue
        match=re.search(r'(-?[0-9]+(?:\.[0-9]+)?)(%)?$',s)
        if not match:entries[str(i)]={'stat':s,'value':'','unit':'unknown','star':None};issues.append('词条'+str(i)+'数字未识别');continue
        label=s[:match.start()].strip(' +:：');unit='percent' if match.group(2) else 'flat'
        legal=model['anchors'].get(anchor,{}).get('pools',{}).get(str(i),[])
        if unit=='percent' and label+'%' in legal:label+='%'
        entries[str(i)]=normalize_stat(label,match.group(1),unit,legal)
        if min(float(x.get('score',0)) for x in xs)<.8:issues.append('词条'+str(i)+'OCR置信度偏低')
    for i,(lo,hi) in enumerate([(385,437),(438,490)],1):
        xs=band(lo,hi);s=text(xs);m=re.search(r'(.+?)[+＋]?([12])$',s)
        zoom=[x for x in items if x.get('zoom') and lo<=x['box'][1]+x['box'][3]/2<=hi]
        z=text(zoom);zm=re.fullmatch(r'[+＋]?([12])',z)
        if zm and all(float(x.get('score',0))>=.8 for x in zoom):
            labels=[x for x in xs if x['box'][0]<450]
            if m and m.group(2)!=zm.group(1) and all(float(x.get('score',0))>=.8 for x in xs):issues.append('特性'+str(i)+'两次OCR结果不一致')
            s=text(labels)+zm.group(1);m=re.search(r'(.+?)([12])$',s)
            xs=labels+zoom
        traits[str(i)]=m.group(1).rstrip('+＋')+m.group(2) if m else None
        if m and any(float(x.get('score',0))<.8 for x in xs):issues.append('特性'+str(i)+'OCR置信度偏低')
        if not m:issues.append('特性'+str(i)+'未识别，需确认是否为空')
    # The image pass adds star evidence; text alone never infers stars from values.
    return {'anchor':anchor,'level':level,'entries':entries,'traits':traits,'issues':issues,'source':'本地OCR，星标未识别'}

def reparse_saved(old,model):
    import vision
    if old.get('rarity') in (1,2):return dict(old)
    obs=parse_ocr(old['raw_ocr'],model);im=vision.image(old['image'])
    obs.update({k:old[k] for k in ('image','original','evidence','raw_ocr') if k in old})
    obs['presence']={};obs['trait_presence']={};obs['source']='已保存截图离线复核'
    for key,item in obs['entries'].items():
        if item:item['star']=vision.star(im,int(key));state='present'
        else:state='absent' if int(key)>=3 and vision.empty_row(im,int(key)) else 'unknown'
        obs['presence'][key]=state
        if state=='absent':obs['issues']=[s for s in obs['issues'] if s!='词条'+key+'未识别，需确认是否为空']
    for key,item in obs['traits'].items():
        state='present' if item else 'absent' if vision.blank_trait(im,int(key)) else 'unknown'
        obs['trait_presence'][key]=state
        if state=='absent':obs['issues']=[s for s in obs['issues'] if s!='特性'+key+'未识别，需确认是否为空']
    return obs
